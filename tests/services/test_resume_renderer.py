import io

import docx

from app.services.resume_renderer import render_cover_letter_docx, render_tailored_resume_docx


def _document(docx_bytes: bytes) -> docx.Document:
    return docx.Document(io.BytesIO(docx_bytes))


def _paragraphs(docx_bytes: bytes) -> list[str]:
    return [paragraph.text for paragraph in _document(docx_bytes).paragraphs]


def _paragraph_by_text(docx_bytes: bytes, text: str):
    for paragraph in _document(docx_bytes).paragraphs:
        if paragraph.text == text:
            return paragraph
    raise AssertionError(f"no paragraph with text {text!r}")


def _full_contact() -> dict[str, str]:
    return {
        "name": "Jane Doe",
        "email": "jane@example.com",
        "phone": "555-1234",
        "location": "Austin, TX",
        "linkedin": "linkedin.com/in/janedoe",
    }


def test_tailored_resume_renders_contact_header():
    docx_bytes = render_tailored_resume_docx(
        "Experienced engineer.", [("Experience", ["Did things."], [])], _full_contact()
    )
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Jane Doe"
    assert paragraphs[1] == "jane@example.com | 555-1234 | Austin, TX | linkedin.com/in/janedoe"
    assert "Experienced engineer." in paragraphs


def test_tailored_resume_omits_header_when_contact_blank():
    docx_bytes = render_tailored_resume_docx("Experienced engineer.", [("Experience", ["Did things."], [])], {})
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Experienced engineer."


def test_tailored_resume_header_defaults_when_contact_omitted():
    docx_bytes = render_tailored_resume_docx("Experienced engineer.", [("Experience", ["Did things."], [])])
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Experienced engineer."


def test_cover_letter_renders_contact_header():
    docx_bytes = render_cover_letter_docx(
        "Dear Hiring Manager,", ["I would be a great fit."], "Sincerely, Jane Doe", _full_contact()
    )
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Jane Doe"
    assert paragraphs[1] == "jane@example.com | 555-1234 | Austin, TX | linkedin.com/in/janedoe"
    assert "Dear Hiring Manager," in paragraphs
    assert "I would be a great fit." in paragraphs
    assert "Sincerely, Jane Doe" in paragraphs


def test_cover_letter_omits_header_when_contact_blank():
    docx_bytes = render_cover_letter_docx(
        "Dear Hiring Manager,", ["I would be a great fit."], "Sincerely,", {}
    )
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Dear Hiring Manager,"


def test_tailored_resume_uppercases_section_headings():
    docx_bytes = render_tailored_resume_docx("Summary.", [("Experience", ["Did things."], [])])
    paragraphs = _paragraphs(docx_bytes)
    assert "EXPERIENCE" in paragraphs
    assert "Experience" not in paragraphs


def test_contact_header_uses_only_available_fields():
    docx_bytes = render_tailored_resume_docx(
        "Summary.",
        [],
        {"name": "Jane Doe", "email": "jane@example.com", "phone": "", "location": "", "linkedin": ""},
    )
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs[0] == "Jane Doe"
    assert paragraphs[1] == "jane@example.com"


def test_entry_renders_bold_title_and_italic_subtitle():
    entries = [
        {
            "title": "Senior Backend Engineer",
            "subtitle": "Acme Corp · Jan 2021 – Present",
            "bullets": ["Led migration to microservices."],
        }
    ]
    docx_bytes = render_tailored_resume_docx("Summary.", [("Experience", [], entries)])
    paragraphs = _paragraphs(docx_bytes)
    assert "Senior Backend Engineer" in paragraphs
    assert "Acme Corp · Jan 2021 – Present" in paragraphs
    assert "Led migration to microservices." in paragraphs

    title_paragraph = _paragraph_by_text(docx_bytes, "Senior Backend Engineer")
    assert title_paragraph.runs[0].bold is True

    subtitle_paragraph = _paragraph_by_text(docx_bytes, "Acme Corp · Jan 2021 – Present")
    assert subtitle_paragraph.runs[0].italic is True


def test_multiple_entries_render_as_separate_blocks():
    entries = [
        {"title": "Senior Engineer", "subtitle": "Acme Corp", "bullets": ["Did A."]},
        {"title": "Engineer", "subtitle": "Widgets Inc", "bullets": ["Did B."]},
    ]
    docx_bytes = render_tailored_resume_docx("Summary.", [("Experience", [], entries)])
    paragraphs = _paragraphs(docx_bytes)
    assert paragraphs.count("Senior Engineer") == 1
    assert paragraphs.count("Engineer") == 1
    assert "Did A." in paragraphs
    assert "Did B." in paragraphs


def test_flat_section_still_renders_without_entries():
    docx_bytes = render_tailored_resume_docx("Summary.", [("Skills", ["Python, Go, Docker"], [])])
    paragraphs = _paragraphs(docx_bytes)
    assert "Python, Go, Docker" in paragraphs


def test_bullet_with_label_prefix_bolds_the_label():
    docx_bytes = render_tailored_resume_docx("Summary.", [("Skills", ["Languages: Python, Go, Java"], [])])
    paragraph = _paragraph_by_text(docx_bytes, "Languages: Python, Go, Java")
    assert paragraph.runs[0].text == "Languages: "
    assert paragraph.runs[0].bold is True
    assert paragraph.runs[1].text == "Python, Go, Java"
    assert not paragraph.runs[1].bold


def test_plain_bullet_without_label_is_not_bolded():
    docx_bytes = render_tailored_resume_docx("Summary.", [("Skills", ["Just a plain skill bullet"], [])])
    paragraph = _paragraph_by_text(docx_bytes, "Just a plain skill bullet")
    assert len(paragraph.runs) == 1
    assert not paragraph.runs[0].bold
